"""A firm's underwriting model, the way one is built by hand, for the template tests."""
import sys, openpyxl
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
wb = openpyxl.Workbook()
s = wb.active; s.title = 'Inputs'
s['A1'] = 'ACME CAPITAL - ACQUISITION MODEL'; s.merge_cells('A1:D1'); s['A1'].font = Font(bold=True, size=14)
labels = {3: 'Property Name', 4: 'Address', 5: 'Purchase Price', 6: 'NOI', 7: 'Cap Rate', 8: 'Units', 9: 'Building SF', 10: 'Occupancy', 12: 'Interest Rate', 13: 'LTV', 14: 'Loan Amount'}
for r, t in labels.items(): s.cell(row=r, column=1, value=t)
s['B5'] = 1000000; s['B5'].number_format = '$#,##0'
s['B6'].number_format = '$#,##0'
s['B7'] = '=IFERROR(B6/B5,"")'; s['B7'].number_format = '0.00%'
s.merge_cells('B10:C10'); s['B10'].number_format = '0.0%'
s['B12'].number_format = '0.00%'; s['B13'].number_format = '0.0%'
s['B14'] = '=IFERROR(B5*B13,"")'
wb.defined_names['InterestRate'] = DefinedName('InterestRate', attr_text='Inputs!$B$12')
r = wb.create_sheet('Rent Roll')
r['A2'] = 'RENT ROLL'
for i, h in enumerate(['Suite', 'Tenant', 'SF', 'Lease Exp', 'Annual Rent', 'Rent / SF'], 1): r.cell(row=4, column=i, value=h)
for i in range(5, 15):
    r.cell(row=i, column=1, value=f'old {i}')
    r.cell(row=i, column=6, value=f'=IFERROR(E{i}/C{i},"")')
sm = wb.create_sheet('Summary')
sm['A1'] = 'Price'; sm['B1'] = '=Inputs!B5'
sm['A2'] = 'Total rent'; sm['B2'] = "=SUM('Rent Roll'!E5:E40)"
sm['A3'] = 'Price / SF'; sm['B3'] = '=IFERROR(Inputs!B5/Inputs!B9,"")'
sm['A4'] = 'Implied cap'; sm['B4'] = '=Inputs!B7'
h = wb.create_sheet('Lists'); h['A1'] = 'helper'; h.sheet_state = 'hidden'
wb.save(sys.argv[1]); print('ok')
