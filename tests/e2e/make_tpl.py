import openpyxl, sys
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Sale Comps'
ws['A1'] = 'ACME REALTY - SALE COMPARABLES'; ws.merge_cells('A1:H1'); ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = 'Prepared for internal use'
hdr = ['Property', 'Address', 'Sale Date', 'Sale Price', 'Building SF', 'Price / SF', 'Cap Rate', 'Notes']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=4, column=i, value=h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7')
# old data that must be cleared, and a $/SF formula column
for r in range(5, 13):
    ws.cell(row=r, column=1, value=f'Old comp {r}'); ws.cell(row=r, column=4, value=1000000 + r); ws.cell(row=r, column=5, value=1000)
    ws.cell(row=r, column=6, value=f'=IFERROR(D{r}/E{r},"")'); ws.cell(row=r, column=4).number_format = '$#,##0'
ws['C14'] = 'Average'; ws['D14'] = '=AVERAGE(D5:D12)'; ws['F14'] = '=IFERROR(SUM(D5:D12)/SUM(E5:E12),"")'
wb.defined_names['CompPrices'] = DefinedName('CompPrices', attr_text="'Sale Comps'!$D$5:$D$12")
s2 = wb.create_sheet('Summary'); s2['A1'] = 'Avg $/SF'; s2['B1'] = "='Sale Comps'!F14"; s2['A2'] = 'Count'; s2['B2'] = "=COUNTA('Sale Comps'!A5:A12)"
s3 = wb.create_sheet('Lists'); s3['A1'] = 'hidden helper'; s3.sheet_state = 'hidden'
wb.save(sys.argv[1]); print('ok')
